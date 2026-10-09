from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent / 'producto_academico'
OUT.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin = sec.right_margin = Cm(2.5)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.size = Pt(12)
doc.styles['Normal'].paragraph_format.line_spacing = 1.5
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
records = []
def p(text, style=None):
    doc.add_paragraph(text, style)
    records.append((text, bool(style)))
def h(text): p(text, 'Heading 1')
def sub(text): p(text, 'Heading 2')
def mark(kind, number, description):
    p(f'(AQUI VA {kind} {number})')
    p(description)
def page(): doc.add_page_break(); records.append(('PAGE', False))

p('[NOMBRE DE LA UNIVERSIDAD]', 'Title')
p('[FACULTAD Y ESCUELA PROFESIONAL]')
p('PRODUCTO ACADÉMICO N.° 2', 'Title')
p('COOPAECA: aplicación web de gestión comercial y contenidos con extensión híbrida para Android', 'Title')
p('Asignatura: [COMPLETAR]')
p('Docente: [COMPLETAR]')
p('Integrantes: [APELLIDOS Y NOMBRES DE CADA INTEGRANTE]')
p('Periodo académico: [COMPLETAR]')
p('Lugar y fecha de presentación: [COMPLETAR]')
p('Nombre de archivo de entrega: PA2_Apellido1_Apellido2_Apellido3')
page()
h('Orientaciones para completar el documento')
p('Este documento desarrolla el proyecto existente y deja espacios para incorporar evidencias académicas. Sustituir los campos entre corchetes y los marcadores antes de entregar. El nombre del archivo debe reemplazar Apellido1, Apellido2 y Apellido3 por los apellidos de los integrantes, según la consigna.')
p('Cada captura debe mostrar la fecha y hora reales del computador. Capturar la ventana de la aplicación junto con la barra de tareas o el panel del reloj del sistema, sin agregar una fecha artificial sobre la imagen. Utilizar datos de prueba y ocultar información personal, contraseñas y claves. Registrar debajo de cada captura el módulo, procedimiento, resultado, fecha y hora.')
p('Las funciones descritas se apoyan en el código del repositorio. Los resultados comprobados se distinguen de las pruebas propuestas. La revisión local no acredita una auditoría integral de seguridad ni la validación de todas las operaciones comerciales. La metodología se presenta como una reconstrucción técnica coherente con la organización del proyecto; el equipo debe ajustarla a su proceso real.')
h('Contenido')
for item in ['1. Nombre del proyecto', '2. Descripción del proyecto', '3. Definición del producto', '4. Arquitectura', '5. Plan de trabajo', '6. Análisis y diseño', '7. Interfaces', '8. Método, metodología y estrategia', '9. Modelo de calidad', '10. Pruebas de validación', '11. Seguridad y contramedidas', '12. Conclusiones', '13. Anexos y referencias']:
    p(item)
page()
h('1. Nombre del proyecto')
p('COOPAECA: sistema web de gestión comercial, administración de contenidos y comunicación institucional, con extensión híbrida para Android.')
h('2. Descripción del proyecto')
p('El proyecto proporciona una plataforma para la Cooperativa Agraria de Exportación y de Calidad Ashaninka. Integra un sitio público para presentar la organización, productos, publicaciones y galerías, y un panel privado para administrar información y operaciones comerciales. Su propósito es centralizar el registro de contenidos, pedidos, ventas, existencias y documentos, reduciendo la dispersión de información y facilitando el seguimiento de las operaciones.')
p('El visitante puede consultar el catálogo y contenido institucional, enviar consultas y registrar un pedido mediante el proceso de compra. Los usuarios autorizados administran los módulos disponibles de acuerdo con sus permisos. El superadministrador dispone de acceso general; otros usuarios reciben acceso explícito por módulo.')
p('Objetivo general: implementar una solución accesible desde navegador y Android que apoye la difusión institucional y el control comercial de COOPAECA, con trazabilidad, validación de datos y controles de acceso.')
p('Objetivos específicos: administrar catálogo y contenidos; registrar y revisar pedidos; gestionar ventas y movimientos de inventario; organizar comprobantes y archivos; controlar usuarios y permisos; y verificar la calidad del producto mediante pruebas documentadas.')
p('Alcance: la solución comprende presentación pública, autenticación, administración, comercio, documentos y comunicación. La aplicación Android consume el servidor web; PHP y MySQL no se ejecutan dentro del teléfono. Las consultas externas de identidad, el correo y los envíos de documentos dependen de sus servicios y configuraciones. Un comprobante generado por el sistema no debe describirse como comprobante electrónico autorizado por SUNAT sin demostrar esa integración.')
h('3. Definición del producto del proyecto')
p('El producto es una aplicación web responsive desarrollada con HTML, CSS y JavaScript, un backend PHP con organización MVC y una base MySQL. Incorpora recursos de PWA —manifiesto, Service Worker y pantalla sin conexión— y un proyecto Android basado en Capacitor. Sus entregables son el código, la base de datos, la documentación y el proyecto móvil; la compilación, firma y funcionamiento del paquete Android deben respaldarse con evidencia independiente.')
sub('3.1. Sustento: aplicativo nativo o híbrido')
p('La versión móvil se clasifica como híbrida porque reutiliza la interfaz web dentro de un contenedor nativo Android provisto por Capacitor. La configuración capacitor.config.json identifica la aplicación como com.ccopaeca.app y dirige su WebView al servidor HTTPS de COOPAECA. La interfaz y la lógica del servidor conservan las tecnologías web, mientras el contenedor permite su distribución en Android. La existencia de archivos Gradle y del proyecto android sustenta esta arquitectura, pero no reemplaza una prueba de instalación del APK.')
p('La versión que se abre en coopaeca.test es la aplicación web. Una PWA instalable y una aplicación híbrida son entregables relacionados, pero distintos: la PWA se instala desde el navegador y el contenedor Capacitor produce el paquete Android. En el repositorio actual, el contenedor apunta al dominio de producción; no apunta a coopaeca.test. Las pruebas móviles requieren un servidor accesible desde el dispositivo.')
p('Esta elección permite reutilizar la interfaz y mantener una lógica de negocio centralizada. Como limitación, las consultas y transacciones necesitan conexión al servidor. La pantalla offline informa la ausencia de conectividad; no demuestra que puedan registrarse ventas o pedidos sin Internet.')
mark('CUADRO', 1, 'Comparación a insertar: aplicación nativa — interfaz y lógica específicas de Android; aplicación web — ejecución en navegador; aplicación híbrida COOPAECA — interfaz web en contenedor Capacitor, backend remoto y acceso HTTPS.')
h('4. Esquema de la arquitectura utilizada')
p('Se utiliza una arquitectura cliente-servidor con separación lógica MVC. El navegador o WebView envía solicitudes HTTP/HTTPS a Apache. public/index.php registra las rutas; Router despacha la solicitud y aplica middleware. Los controladores coordinan la operación, los servicios implementan reglas de negocio y Database proporciona la conexión PDO. Las vistas generan la respuesta y MySQL almacena los datos.')
p('Flujo representativo: usuario → navegador/WebView → servidor Apache → ruta y middleware → controlador → servicio → PDO/MySQL → vista o respuesta JSON → usuario. El sistema combina páginas renderizadas en el servidor y consultas JSON; no se presenta como una API REST completa.')
p('Capas: presentación en app/Views y public/assets; entrada y control en public/index.php, app/Core y app/Controllers; negocio en app/Services; persistencia mediante PDO, SQL y el esquema de database. La presencia de una clase base Model no implica que todas las entidades tengan modelos propios: varios controladores y servicios consultan directamente mediante PDO.')
mark('GRAFICO', 1, 'Arquitectura: dibujar los clientes navegador y Android/Capacitor, conexión HTTPS de producción, Apache/PHP, Router/middleware, controladores, servicios, vistas, PDO, MySQL y almacenamiento de archivos. Añadir servicios externos IMAP/SMTP e identidad como dependencias externas. Diferenciar el entorno local HTTP del entorno productivo HTTPS.')
p('Entorno local: carpeta C:/laragon/www/ccopaeca; DocumentRoot dirigido a public; URL prevista http://coopaeca.test; base restaurada coopaeca_local_20261007. La comprobación directa del dominio falló por falta de resolución local; forzando el dominio a 127.0.0.1, inicio, login y ubigeo respondieron HTTP 200. El registro en hosts sigue siendo necesario para el acceso normal mientras no se complete ese ajuste.')
h('5. Plan de trabajo')
p('Se propone un cronograma de seis semanas. Es un plan de organización académica y cierre del proyecto, no una afirmación de fechas históricas. Incorporar responsables y fechas reales del equipo.')
phases = [
('Semana 1', 'Requisitos y alcance', 'Identificar actores, operaciones y reglas; revisar módulos y dependencias.', 'Lista de requisitos y criterios de aceptación.'),
('Semana 2', 'Análisis y diseño', 'Modelar casos de uso, clases conceptuales, navegación y estructura de datos.', 'Diagramas revisados y diseño de interfaces.'),
('Semana 3', 'Implementación del sitio y acceso', 'Revisar sitio público, contenidos, autenticación y permisos.', 'Módulos accesibles y validaciones básicas.'),
('Semana 4', 'Implementación comercial', 'Revisar pedidos, ventas, inventario, pagos y documentos.', 'Flujos integrados y consistencia de datos.'),
('Semana 5', 'Calidad y seguridad', 'Ejecutar pruebas unitarias, integración, permisos, archivos y compatibilidad.', 'Matriz de resultados e incidencias corregidas.'),
('Semana 6', 'Móvil y entrega', 'Validar PWA/Android, elaborar capturas con reloj y completar anexos.', 'Informe final y evidencias reproducibles.')]
for week, phase, action, deliverable in phases: p(f'{week} — {phase}. Actividades: {action} Entregable: {deliverable} Responsable: [COMPLETAR].')
mark('CUADRO', 2, 'Cronograma: columnas semana/fecha, actividad, responsable, dependencia, entregable y estado. Convertir las seis fases anteriores en tabla o diagrama de Gantt.')
p('Criterio de cierre: cada requisito prioritario debe vincularse con una prueba, cada incidencia crítica debe corregirse y comprobarse nuevamente, y cada módulo debe disponer de código y captura de evidencia. Mantener copia de seguridad antes de alterar datos o configuración.')
h('6. Presentación del análisis y diseño')
sub('6.1. Actores y requisitos')
p('Actores: visitante/cliente, usuario administrativo autorizado y superadministrador. Los servicios de identidad y correo son sistemas externos. Un cliente puede realizar un pedido desde la parte pública; no se presume una cuenta de cliente si el flujo implementado no la exige.')
p('Requisitos funcionales: RF01 consultar información pública; RF02 iniciar/cerrar sesión y recuperar acceso; RF03 administrar usuarios y permisos; RF04 mantener contenidos y archivos; RF05 administrar productos y medios de pago; RF06 registrar/revisar pedidos; RF07 registrar/cancelar ventas y consultar comprobantes; RF08 gestionar movimientos y existencias; RF09 atender contactos y correo; RF10 modificar configuración institucional.')
p('Requisitos no funcionales: controlar acceso en el servidor; mantener integridad de operaciones; adaptar interfaces a diferentes pantallas; proteger datos sensibles; registrar errores; conservar trazabilidad; y documentar restauración y despliegue. Los objetivos medibles se definen en el apartado de calidad y se validan antes de declararlos alcanzados.')
sub('6.2. Casos de uso y clases por módulo')
p('En los diagramas UML de casos de uso, ubicar a los actores fuera del límite del sistema y asociarlos a sus acciones. Utilizar «include» para comportamiento obligatorio reutilizado y «extend» únicamente para comportamiento opcional con condición. Las verificaciones de sesión y permisos pueden indicarse como precondiciones.')
p('Los diagramas de clases siguientes son modelos de análisis y diseño: las entidades representan conceptos y tablas, no necesariamente clases PHP ya implementadas. Diferenciar entidades, controladores y servicios. Para el código real usar los nombres del repositorio y evitar inventar métodos. Representar asociaciones, atributos principales y multiplicidades respaldadas por el esquema.')
modules = [
('Sitio público y compra', 'Visitante/cliente', 'Consultar inicio, nosotros, catálogo, publicaciones y galería; buscar contenido; gestionar compra y registrar pedido; enviar consulta.', 'LandingController, CheckoutController, Producto, Pedido, DetallePedido y MensajeContacto. Pedido 1 → 1..* DetallePedido; Producto 1 → 0..* DetallePedido.', 'app/Controllers/LandingController.php; app/Controllers/CheckoutController.php', 'Abrir catálogo, seleccionar producto, registrar pedido de prueba y comprobar confirmación.'),
('Autenticación y perfil', 'Usuario y superadministrador', 'Iniciar sesión, cerrar sesión, solicitar recuperación y actualizar perfil.', 'AuthController, ProfileController, Usuario, Rol, TokenRecordatorio y RecuperaciónAcceso. Rol 1 → 0..* Usuario.', 'app/Controllers/AuthController.php; app/Controllers/ProfileController.php; app/Helpers/auth.php', 'Probar credenciales válidas e inválidas y comprobar que una ruta privada no abre sin sesión.'),
('Usuarios y permisos', 'Superadministrador o usuario autorizado', 'Listar, crear, editar y activar/desactivar usuarios; asignar acceso por módulo.', 'UserController, Usuario, Rol y AccesoModulo. Usuario 1 → 0..* AccesoModulo; el permiso efectivo se evalúa por usuario.', 'app/Controllers/UserController.php; app/Helpers/auth.php', 'Asignar acceso a un módulo a una cuenta de prueba y contrastar acceso permitido y denegado.'),
('Productos y medios de pago', 'Usuario autorizado', 'Mantener productos, imágenes y categorías; configurar métodos de pago, cuenta y moneda.', 'ProductController, PaymentMethodController, Producto, Categoria, Archivo y MetodoPago. Producto y Categoria presentan relación muchos a muchos mediante product_category.', 'app/Controllers/ProductController.php; app/Controllers/PaymentMethodController.php', 'Registrar y editar producto de prueba; validar precio, datos obligatorios y presentación del medio de pago.'),
('Pedidos', 'Cliente y usuario autorizado', 'Registrar pedido desde compra; listar, revisar, aprobar o rechazar; consultar documentos y gestionar envío.', 'OrderController, OrderService, Pedido, DetallePedido, Producto y Archivo. Relacionar el pedido con sus detalles y su documentación; comprobar la relación con ventas en el esquema.', 'app/Controllers/OrderController.php; app/Services/OrderService.php', 'Registrar pedido, revisar detalle y probar aprobación/rechazo en una copia de pruebas; verificar estado e inventario.'),
('Ventas y comprobantes', 'Usuario autorizado', 'Registrar venta, consultar detalle, cancelar venta y emitir/consultar comprobante.', 'SalesController, SaleService, ReceiptService, Venta, DetalleVenta y Archivo. Venta 1 → 1..* DetalleVenta; Producto 1 → 0..* DetalleVenta.', 'app/Controllers/SalesController.php; app/Services/SaleService.php; app/Services/ReceiptService.php', 'Crear venta con stock disponible, verificar total y documento; cancelar y comprobar restitución conforme a la regla implementada.'),
('Inventario', 'Usuario autorizado', 'Consultar existencias y movimientos; registrar operaciones de inventario según las acciones disponibles.', 'InventoryController, InventoryService, Producto y MovimientoStock. Producto 1 → 0..* MovimientoStock; las operaciones comerciales interactúan con InventoryService.', 'app/Controllers/InventoryController.php; app/Services/InventoryService.php', 'Comparar existencias antes/después de una operación; intentar cantidad superior al disponible y comprobar consistencia.'),
('Publicaciones y servicios', 'Editor autorizado y visitante', 'Crear, editar y eliminar publicaciones y servicios; mostrar contenidos públicos.', 'PostController, ServiceController, Publicacion, Categoria, Etiqueta y Servicio. Publicacion y Categoria se relacionan mediante post_category.', 'app/Controllers/PostController.php; app/Controllers/ServiceController.php', 'Publicar contenido de prueba, comprobar visualización pública y validar entradas inválidas.'),
('Galerías, sliders y archivos', 'Editor autorizado y visitante', 'Administrar galerías, imágenes, sliders y archivos; visualizar recursos públicos.', 'GalleryController, SliderController, MediaController, Galeria, ImagenGaleria, Slider, ElementoSlider y Archivo. Galeria 1 → 0..* ImagenGaleria; Slider 1 → 0..* ElementoSlider.', 'app/Controllers/GalleryController.php; app/Controllers/SliderController.php; app/Controllers/MediaController.php', 'Subir imagen permitida, comprobar presentación e intentar archivo no permitido en el entorno de pruebas.'),
('Contactos y correo', 'Visitante y usuario autorizado; servidor de correo externo', 'Enviar contacto, revisar mensajes, administrar cuenta de correo, consultar buzón, redactar y gestionar borradores.', 'ContactController, MailboxController, MailAccountController, MailComposeController, ImapService, SmtpService, MensajeContacto, CuentaCorreo y Borrador. CuentaCorreo se asocia con datos de correo según las claves reales del esquema.', 'app/Controllers/ContactController.php; app/Controllers/Mail/; app/Services/ImapService.php; app/Services/SmtpService.php', 'Registrar consulta local y comprobar su almacenamiento; probar correo únicamente con cuenta de pruebas y servicio autorizado.'),
('Configuración y presentación institucional', 'Usuario autorizado', 'Actualizar configuración, información institucional y redes sociales; consultar indicadores del panel.', 'SettingsController, AboutController, SocialNetworkController, DashboardController, Configuracion y RedSocial. La configuración se representa como claves/valores.', 'app/Controllers/SettingsController.php; app/Controllers/AboutController.php; app/Controllers/SocialNetworkController.php; app/Controllers/DashboardController.php', 'Cambiar un valor en la copia de pruebas y comprobar su efecto en la interfaz; verificar restricciones de acceso.')]
for i,(name,actors,uc,classes,code,test) in enumerate(modules, 1):
    sub(f'6.2.{i}. {name}')
    p(f'Actores: {actors}. Casos de uso: {uc}')
    p('Precondición: sesión y permisos cuando corresponda. Resultado esperado: operación autorizada, datos válidos y respuesta comprensible. Ante error, conservar consistencia y mostrar un mensaje sin exponer información sensible.')
    mark('GRAFICO', i*2, f'Diagrama de casos de uso del módulo {name}: representar los actores y las acciones descritas.')
    p(f'Clases y asociaciones propuestas: {classes}')
    mark('GRAFICO', i*2+1, f'Diagrama de clases del módulo {name}: incluir atributos, asociaciones y multiplicidades; distinguir entidades conceptuales de clases existentes.')
    p(f'Referencia de implementación: {code}.')
sub('6.3. Mapa de navegación web')
p('Navegación pública: Inicio → Nosotros / Producto / Publicación / Galería / Buscar / Compra. El acceso a administración parte de Login y conduce a Dashboard. Las pantallas privadas incluyen perfil, usuarios, productos, pedidos, ventas, inventario, medios de pago, publicaciones, servicios, galerías, sliders, medios, contactos, correo, configuración, información institucional y redes sociales; su visibilidad depende de los permisos. Cerrar sesión devuelve al usuario al contexto público.')
mark('GRAFICO', 24, 'Mapa de navegación: separar área pública y privada. Mostrar login como entrada al panel; anotar permisos y relacionar listados con formularios de creación, edición y detalle. Contrastar las rutas exactas con public/index.php.')
sub('6.4. Diseño de datos y reglas de negocio')
p('Las entidades comerciales principales son products, orders, order_items, sales, sale_items y stock_movements. Los archivos se registran en files; usuarios y accesos en users, roles y user_module_access. Las claves foráneas vinculan registros y las transacciones coordinan cambios relacionados. InventoryService utiliza SELECT ... FOR UPDATE para bloquear registros durante operaciones sensibles; los servicios comerciales abren transacciones. Esto aporta un mecanismo para controlar concurrencia, pero su funcionamiento completo debe probarse con operaciones simultáneas.')
mark('GRAFICO', 25, 'Diagrama entidad-relación: productos, pedidos/detalles, ventas/detalles, movimientos, usuarios y archivos. Obtener las claves y cardinalidades exactas de dashboard_base.sql. Omitir datos personales del volcado.')
h('7. Explicación e ilustración de las interfaces')
p('La interfaz pública facilita la consulta institucional y comercial; el panel organiza acciones administrativas mediante navegación lateral, listados y formularios. Las siguientes imágenes deben sustituirse por capturas reales. En cada una mantener visible la fecha y hora del computador y explicar qué acción puede realizar el usuario.')
interfaces = [
('Inicio público', 'Presenta identidad institucional, contenido destacado y accesos al catálogo. Capturar la cabecera y una sección principal.'),
('Nosotros', 'Presenta información de la organización y su identidad. Capturar el contenido institucional disponible.'),
('Catálogo y detalle de producto', 'Muestra productos, imágenes e información comercial. Capturar listado y un detalle con datos de prueba.'),
('Publicaciones y galería', 'Permite consultar noticias y recursos visuales. Capturar una publicación y una galería.'),
('Compra y confirmación', 'Recoge información del pedido y muestra el resultado del registro. Capturar formulario y confirmación, sin datos personales reales.'),
('Login y recuperación', 'Permite autenticación y recuperación de acceso. Capturar formulario y un mensaje de validación, sin mostrar contraseñas.'),
('Dashboard y perfil', 'Presenta el área privada y edición del perfil. Capturar indicadores visibles y formulario de perfil con datos anonimizados.'),
('Usuarios y permisos', 'Gestiona usuarios y accesos por módulo. Capturar listado y asignación de permisos con cuenta de prueba.'),
('Productos y medios de pago', 'Gestiona catálogo y formas de pago. Capturar listado, formulario y configuración con números de cuenta ficticios.'),
('Pedidos', 'Organiza revisión, detalle y acciones autorizadas. Capturar un pedido de prueba y su estado.'),
('Ventas y comprobantes', 'Muestra ventas, detalle y documentos. Capturar una venta y un comprobante de prueba, sin atribuir validez fiscal no demostrada.'),
('Inventario', 'Presenta existencias y movimientos. Capturar producto y movimientos que permitan explicar la trazabilidad.'),
('Contenidos y recursos', 'Permite administrar publicaciones, servicios, galerías, sliders y archivos. Agregar una captura por pantalla implementada.'),
('Contactos y correo', 'Centraliza consultas y operaciones de correo. Capturar datos de prueba; ocultar credenciales y comunicaciones privadas.'),
('Configuración institucional', 'Permite actualizar parámetros, información y redes sociales. Capturar las opciones del proyecto sin exponer secretos.')]
for i,(name,description) in enumerate(interfaces,1):
    sub(f'7.{i}. {name}'); p(description)
    mark('IMAGEN', i, f'Captura de {name}. Fecha/hora visible: [COMPLETAR]. Explicación de la acción mostrada: [COMPLETAR].')
h('8. Método, metodología y estrategia de implementación')
sub('8.1. Método')
p('Se plantea un método de desarrollo aplicado: identificar el problema de gestión, analizar los procesos, diseñar una solución, implementarla y verificarla mediante evidencia reproducible. Para cada requisito se establece una operación esperada, una regla de negocio y un criterio de aceptación. Esta secuencia permite relacionar las necesidades de la cooperativa con decisiones concretas del sistema.')
sub('8.2. Metodología')
p('La organización modular y el avance por funcionalidades son compatibles con una metodología iterativa e incremental. Cada incremento comprende análisis, diseño, implementación, prueba y ajuste; por ejemplo, completar autenticación antes de habilitar administración y completar productos antes de integrar ventas e inventario. Esta descripción debe confirmarse con el historial real del equipo. No se afirma que se aplicó Scrum formal porque no se han acreditado sus roles, eventos y artefactos.')
sub('8.3. Estrategia')
p('La estrategia prioriza componentes comunes —conexión, rutas, sesiones, permisos y vistas— y posteriormente integra módulos. Se reutiliza la interfaz web en Android mediante Capacitor; se mantiene una copia local de datos para verificar cambios; se emplea Git para controlar versiones y se revisan dependencias con Composer. Las pruebas deben acompañar cada incremento y enfocarse primero en permisos, integridad comercial y archivos.')
mark('CUADRO', 3, 'Trazabilidad a insertar: requisito → módulo → archivo/controlador → prueba → resultado → captura. Ejemplo: RF02 → autenticación → AuthController.php → acceso inválido → pendiente → imagen de evidencia.')
p('Evidencias metodológicas a completar: historial de cambios, lista de tareas, decisiones de diseño, registro de incidencias y revisiones del equipo. La existencia del repositorio aporta control de versiones; no demuestra por sí misma las reuniones o validaciones con usuarios.')
h('9. Modelo de calidad del proyecto')
p('Se propone ISO/IEC 25010:2023 como marco para organizar la evaluación de calidad del producto. Se utiliza como referencia para establecer criterios verificables; no se declara certificación ni cumplimiento integral. Fuente oficial: https://www.iso.org/standard/78176.html. Las métricas siguientes son objetivos propuestos para COOPAECA, no resultados ya obtenidos.')
quality = [
('Adecuación funcional', 'Cubrir consulta, administración y operaciones comerciales.', 'Aprobar el 100 % de casos críticos priorizados antes de la entrega.'),
('Eficiencia de desempeño', 'Responder adecuadamente bajo una carga definida.', 'Objetivo inicial: percentil 95 menor o igual a 3 s con 10 usuarios concurrentes; documentar equipo, datos y escenario.'),
('Compatibilidad', 'Mantener comportamiento entre navegadores y servicios.', 'Validar flujos prioritarios en Chrome y Edge y en el WebView Android usado en pruebas.'),
('Capacidad de interacción', 'Facilitar comprensión y ejecución de tareas.', 'Tres usuarios de prueba completan consulta y registro sin errores bloqueantes; comprobar etiquetas, mensajes y navegación por teclado.'),
('Fiabilidad', 'Conservar integridad ante errores y concurrencia.', 'No producir operaciones parciales ni stock incoherente en los escenarios definidos de fallo y simultaneidad.'),
('Seguridad', 'Proteger acceso, datos y archivos.', 'Aprobar casos de permisos, CSRF, consultas parametrizadas y archivos; corregir hallazgos críticos identificados.'),
('Mantenibilidad', 'Facilitar comprensión y cambios.', 'Vincular módulos con archivos, documentar instalación y ejecutar las pruebas existentes tras cada cambio relevante.'),
('Flexibilidad', 'Adaptarse a entorno y tamaño de pantalla.', 'Validar pantallas principales en escritorio y móvil y documentar despliegue en otra instalación.'),
('Safety: prevención del daño', 'Evitar efectos perjudiciales de operaciones incorrectas.', 'Comprobar validaciones de cantidades y totales, confirmaciones de operaciones sensibles y capacidad de recuperar datos.')]
for name, aim, criterion in quality: p(f'{name}: {aim} Criterio: {criterion}')
mark('CUADRO', 4, 'Matriz de calidad: característica, criterio, métrica, instrumento, responsable, resultado y evidencia. Mantener en pendiente los objetivos aún no medidos.')
h('10. Pruebas de validación del proyecto')
sub('10.1. Estrategia y tipos de pruebas')
p('Las pruebas se realizarán sobre una copia controlada de la base, con usuarios de prueba y datos ficticios. Incluir casos positivos, negativos, valores límite y operaciones repetidas. Registrar versión del código, ambiente, datos, pasos, resultado esperado, resultado observado y evidencia. No ejecutar escenarios de carga o seguridad sobre servicios externos sin autorización.')
p('Pruebas unitarias: verifican unidades concretas; existen pruebas de formato de códigos de pedidos/ventas y selección de idioma. Pruebas de integración: verifican controlador, servicio y base conjuntamente. Pruebas de sistema: recorren procesos completos. Pruebas de aceptación: un usuario del negocio confirma que la operación responde a su necesidad. Pruebas de regresión: comprueban que un cambio no rompe funcionalidades previamente validadas.')
p('Pruebas adicionales propuestas: adaptación a pantallas y navegadores, carga controlada, simultaneidad de inventario, restauración de respaldo, permisos y archivos, instalación PWA/Android y recuperación al reconectar. El tipo de prueba se elige por el riesgo: una venta requiere consistencia transaccional; una página pública requiere presentación y navegación; el acceso requiere verificación de autenticación y autorización.')
sub('10.2. Casos de prueba propuestos por módulo')
for i,(name,actors,uc,classes,code,test) in enumerate(modules,1):
    p(f'PV-{i:02d} — {name}. Procedimiento: {test} Resultado esperado: la función válida se completa y la entrada inválida o no autorizada se rechaza sin alterar datos indebidamente. Estado: pendiente de ejecución y captura formal.')
extra = [
('PV-12', 'CSRF', 'Enviar una operación POST sin token y con token inválido.', 'Rechazo 419 y ausencia de modificación.'),
('PV-13', 'Fuerza bruta', 'Repetir intentos inválidos de acceso en la misma sesión.', 'Rechazo 429 al superar el límite; repetir con sesión nueva para documentar la limitación del control actual.'),
('PV-14', 'XSS', 'Guardar texto con etiquetas de script en campos de prueba y revisar salida.', 'No ejecutar scripts; evaluar cada contexto de salida y contenido enriquecido.'),
('PV-15', 'Inyección SQL', 'Introducir cadenas especiales como datos en formularios de prueba.', 'No alterar la consulta ni revelar SQL; rechazo o tratamiento como valor literal.'),
('PV-16', 'Concurrencia', 'Dos sesiones intentan vender el último stock disponible.', 'Solo se completan operaciones compatibles con la existencia; no hay stock negativo.'),
('PV-17', 'Integridad y rollback', 'Provocar un fallo controlado durante una operación de prueba.', 'No conservar venta o movimiento parcial; documentar el mecanismo utilizado.'),
('PV-18', 'PWA y Android', 'Instalar, navegar, usar botón Atrás, desconectar y reconectar.', 'Interfaz usable y aviso offline; las operaciones remotas no simulan éxito sin conexión.'),
('PV-19', 'Restauración', 'Restaurar respaldo en una base independiente y comparar registros.', 'Datos y relaciones recuperados y aplicación conectada correctamente.'),
('PV-20', 'Desempeño', 'Ejecutar carga controlada según el objetivo del apartado 9.', 'Registrar percentiles, errores y recursos; comparar con el objetivo propuesto.')]
for id,name,steps,expected in extra: p(f'{id} — {name}. Pasos: {steps} Resultado esperado: {expected} Estado: propuesto, pendiente.')
mark('CUADRO', 5, 'Matriz de ejecución: ID, requisito/módulo, precondiciones, datos, pasos, esperado, observado, aprobado/fallido, fecha/hora y evidencia. No rellenar observado ni aprobado antes de ejecutar.')
sub('10.3. Resultados realmente comprobados en el entorno local')
p('En la revisión local se ejecutó PHPUnit: cuatro pruebas y cuatro aserciones correctas. Su alcance se limita a formato de códigos y traducciones; no valida por sí mismo ventas, pagos ni seguridad integral. Las dependencias instaladas superaron la comprobación de requisitos de plataforma de Composer.')
p('La restauración de dashboard_base.sql produjo 45 tablas y 60 relaciones de clave foránea. Se encontraron 2 usuarios, 6 productos, 18 pedidos, 17 ventas y 1892 registros de ubigeo. Los 75 registros de archivos tenían su archivo físico. Estas cantidades describen la copia revisada, no un objetivo funcional del sistema.')
p('Inicio, login y consulta de departamentos respondieron HTTP 200 al forzar coopaeca.test hacia 127.0.0.1. La consulta de ubigeo devolvió una respuesta JSON válida. Una sesión local de prueba permitió revisar dashboard, productos y configuración; pedidos, ventas e inventario devolvieron 403 para el usuario admin seleccionado, coherente con su acceso restringido. Esto no reemplaza la ejecución de esos módulos con un usuario autorizado.')
p('La resolución normal de coopaeca.test no funcionó en la última revisión por falta de registro en hosts. Debe corregirse y repetirse la prueba desde el navegador. No se tomaron capturas formales de estas verificaciones: se incorporarán con el reloj visible. Tampoco se acreditaron envío de correo, integraciones externas, prueba de carga ni instalación Android.')
mark('IMAGEN', 16, 'Captura de PHPUnit con cuatro pruebas correctas y fecha/hora visible. Ejecutar: php vendor/bin/phpunit --configuration phpunit.xml.')
mark('IMAGEN', 17, 'Captura de navegación en coopaeca.test después de registrar el dominio, con fecha/hora visible. Mostrar inicio y acceso al panel.')
mark('IMAGEN', 18, 'Captura de una operación permitida y otra denegada por permisos, indicando el tipo de usuario de prueba.')
h('11. Medidas y contramedidas de seguridad')
p('El objetivo es reducir la probabilidad y el impacto de ataques. Ninguna lista de controles garantiza vulnerabilidad cero. Se documentan medidas presentes en el código y mejoras propuestas; su eficacia debe verificarse mediante pruebas. OWASP ASVS se propone como referencia de requisitos y verificación: https://owasp.org/projects/asvs.')
sub('11.1. Medidas presentes en la implementación')
controls = [
('Autenticación y contraseñas', 'AuthController utiliza password_verify y password_hash; se evita almacenar la contraseña en texto plano. La sesión se regenera al iniciar sesión. Verificar política de contraseñas y recuperación mediante pruebas.'),
('Autorización en el servidor', 'AuthMiddleware y PermissionMiddleware protegen rutas. El superadministrador tiene acceso general y otros usuarios se evalúan mediante user_module_access. Ocultar opciones del menú no reemplaza la comprobación del servidor.'),
('Consultas parametrizadas', 'La conexión PDO desactiva la emulación de preparaciones y el código utiliza prepare/execute para valores. Los identificadores SQL dinámicos, como columnas u ordenamientos, requieren listas permitidas; no se afirma que PDO asegure automáticamente cualquier consulta.'),
('Protección CSRF', 'Se genera un token aleatorio de sesión y se compara con hash_equals. Las solicitudes modificadoras incorporan validación mediante middleware.'),
('Sesiones y cookies', 'Se configuran HttpOnly, SameSite=Lax, modo estricto y cookies de sesión. El atributo Secure depende de APP_SECURE_COOKIES; en el entorno HTTP local está desactivado y debe habilitarse junto con HTTPS en producción.'),
('Salida y contenido', 'La aplicación dispone del helper e() para escapar HTML y de la dependencia HTMLPurifier. Debe comprobarse el uso efectivo de escape o sanitización en cada campo y contexto; instalar una biblioteca no prueba protección completa frente a XSS.'),
('Limitación de intentos', 'RateLimitMiddleware limita acceso y otras operaciones; login usa cinco intentos por 60 segundos. El registro de intentos reside en la sesión y combina una clave con la IP: iniciar otra sesión puede eludir esa memoria, por lo que no es un bloqueo global por IP.'),
('Cifrado de secretos de correo', 'app/Helpers/crypto.php utiliza libsodium y una clave definida en el entorno para cifrar credenciales de cuentas de correo. Esta clave no debe incluirse en el informe; su protección y respaldo son necesarios.'),
('Documentos y archivos', 'SecureDocumentService comprueba rutas y tipo MIME de documentos privados. Las reglas .htaccess protegen ubicaciones sensibles y los directorios de documentos. Deben validarse también los controladores de subida y las comprobaciones de propiedad/permisos por documento.'),
('Configuración del servidor', 'El DocumentRoot local apunta a public. Se deshabilita el listado de directorios y se incorporan cabeceras como X-Frame-Options y X-Content-Type-Options. Los secretos de .env permanecen fuera del directorio público.'),
('Transacciones y trazabilidad', 'Los servicios comerciales utilizan transacciones y bloqueos de filas para mantener consistencia; el sistema dispone de registros de actividad y errores. Se deben verificar cobertura de auditoría y ausencia de información sensible en logs.'),
('Cliente híbrido y caché', 'Capacitor configura servidor HTTPS, cleartext=false, allowMixedContent=false y depuración WebView desactivada. La PWA limita la caché a recursos estáticos; revisar Service Worker y acceso offline como parte de las pruebas.')]
for name,detail in controls: p(f'{name}: {detail}')
sub('11.2. Contramedidas adicionales propuestas')
counter = [
('Prioridad alta — despliegue', 'Forzar HTTPS, activar cookies Secure y desactivar APP_DEBUG en producción; verificar certificado y errores genéricos. Evidencia: inspección de configuración, cookies y respuesta ante fallo controlado.'),
('Prioridad alta — privilegios de base', 'Crear usuario MySQL exclusivo del sistema con los privilegios mínimos. La cuenta local actual dispone de permisos amplios y no debe reutilizarse en producción. Evidencia: revisar permisos del usuario dedicado.'),
('Prioridad alta — ataques al acceso', 'Persistir límites en base de datos o caché del servidor; combinar cuenta e IP, aplicar espera progresiva y alertas; incorporar MFA para cuentas privilegiadas. Evidencia: intentos desde sesiones distintas y acceso con segundo factor.'),
('Prioridad alta — secretos y archivos', 'Rotar secretos si fueron expuestos, restringir lectura de .env y mantener credenciales fuera de repositorios y anexos. Endurecer subidas con tamaño máximo, tipos permitidos, nombres aleatorios y bloqueo de ejecución. Evidencia: archivos válidos/ilegales y revisión de permisos.'),
('Prioridad alta — recuperación', 'Programar respaldos cifrados de base y archivos, guardar copia separada y realizar restauración de ensayo. Definir responsables, frecuencia y objetivos de recuperación. Evidencia: restauración y comparación de registros.'),
('Prioridad media — XSS y superficie web', 'Revisar contextos de escape, sanitización de HTML y enlaces; incorporar Content-Security-Policy tras evaluar recursos utilizados. Evidencia: casos XSS controlados y revisión de cabeceras.'),
('Prioridad media — dependencias', 'Auditar Composer y npm, evaluar avisos y actualizar componentes en una rama con pruebas. Evidencia: reporte de auditoría y regresión posterior; la comprobación de plataforma no sustituye la auditoría.'),
('Prioridad media — monitoreo e incidentes', 'Centralizar logs, alertar intentos repetidos y operaciones sensibles, definir procedimiento de contención, recuperación y revisión de causas. Evidencia: simulación controlada de alerta y registro de respuesta.'),
('Prioridad media — documentos y concurrencia', 'Comprobar autorización por objeto en toda descarga y prevenir duplicación de operaciones sensibles mediante idempotencia cuando corresponda. Evidencia: cambio de identificadores, doble envío y operación simultánea.'),
('Prioridad media — Android', 'Validar navegación permitida, enlaces externos, firma del paquete y configuración final de WebView; mantener el dispositivo y WebView actualizados. Evidencia: revisión de paquete y pruebas en teléfono.')]
for name,detail in counter: p(f'{name}: {detail}')
mark('CUADRO', 6, 'Matriz de riesgos: amenaza, activo, probabilidad, impacto, control actual, limitación, contramedida, responsable y prueba. Separar expresamente controles implementados de mejoras pendientes.')
h('12. Conclusiones')
p('1. COOPAECA integra difusión institucional y administración comercial mediante una aplicación web. Su extensión Android basada en Capacitor sustenta la clasificación híbrida de la versión móvil, mientras la ejecución en coopaeca.test corresponde al componente web.')
p('2. La arquitectura cliente-servidor y la separación de rutas, controladores, servicios, vistas y persistencia permiten organizar la solución por módulos. La integridad de inventario, pedidos y ventas necesita validación transaccional y de concurrencia, además de revisar pantallas.')
p('3. La copia importada y las verificaciones locales aportan evidencia inicial de funcionamiento. Las cuatro pruebas unitarias tienen alcance limitado; faltan evidencias por módulo, aceptación de usuarios y pruebas de Android, integraciones y desempeño.')
p('4. Las medidas de autenticación, autorización, CSRF, consultas parametrizadas, sesiones y cifrado reducen riesgos, pero requieren comprobación sistemática. Se priorizan HTTPS productivo, mínimos privilegios, límites persistentes, respaldo y protección de secretos.')
p('5. El modelo de calidad propuesto vincula requisitos con métricas y evidencias. La entrega académica debe completar diagramas, código y capturas con fecha y hora reales, y ajustar la metodología y cronograma al trabajo realizado por el equipo.')
h('13. Anexos')
sub('Anexo A. Código de implementación de cada módulo')
p('Incorporar código real o extractos representativos con ruta de origen y explicación de su función. Si la consigna exige código completo, adjuntar los archivos de cada módulo además de los extractos. No incluir .env, claves, contraseñas, dumps con datos personales ni credenciales de correo.')
for i,(name,actors,uc,classes,code,test) in enumerate(modules,1):
    p(f'A.{i}. {name}')
    p(f'(AQUI VA CODIGO DEL MODULO {name.upper()})')
    p(f'Archivos base: {code}. Incluir vistas relacionadas, registro de rutas y validaciones pertinentes. Explicar entradas, reglas, persistencia y respuesta.')
p('(AQUI VA CODIGO DEL CONTENEDOR ANDROID Y LA PWA)')
p('Incluir capacitor.config.json, public/manifest.webmanifest, public/sw.js y public/assets/js/pwa.js. Explicar servidor remoto, controles de WebView, alcance de caché y dependencia de conectividad.')
sub('Anexo B. Código de conexión a la base de datos')
p('(AQUI VA CODIGO DE CONEXION A LA BASE DE DATOS)')
p('Incluir app/Core/Database.php y app/Config/database.php. Explicar DSN MySQL, charset utf8mb4, credenciales mediante variables de entorno, PDO::ERRMODE_EXCEPTION, FETCH_ASSOC y ATTR_EMULATE_PREPARES=false. No copiar valores secretos; usar ejemplos ficticios.')
p('(AQUI VA CODIGO DE UNA CONSULTA PARAMETRIZADA Y UNA TRANSACCION)')
p('Seleccionar ejemplos reales de OrderService, SaleService o InventoryService; explicar prepare/execute, beginTransaction, commit/rollback y bloqueo FOR UPDATE. Documentar por qué deben ejecutarse como una unidad consistente.')
sub('Anexo C. Capturas de pruebas por módulo')
p('Todas las pruebas siguientes requieren fecha y hora del computador visibles. No presentar marcadores como evidencia completada. Si una prueba falla, registrar el fallo, la corrección y la nueva ejecución.')
for i,(name,actors,uc,classes,code,test) in enumerate(modules,1):
    mark('IMAGEN', 18+i, f'Evidencia PV-{i:02d}, módulo {name}: {test} Registrar esperado, observado, estado, fecha/hora y datos ficticios usados.')
mark('IMAGEN', 30, 'Pruebas de seguridad: CSRF, permisos, límite de intentos y archivos no permitidos; agregar capturas separadas identificadas por ID de prueba.')
mark('IMAGEN', 31, 'Pruebas de integridad: concurrencia, rollback y restauración; mostrar resultados relevantes con reloj visible.')
mark('IMAGEN', 32, 'Pruebas Android/PWA: instalación, navegación y ausencia de conexión. Mantener fecha/hora visible en el teléfono y/o computador según la exigencia del docente.')
sub('Anexo D. Ficha de registro de evidencia')
for label in ['ID de prueba', 'Módulo y requisito', 'Versión y ambiente', 'Usuario de prueba y permisos', 'Precondiciones y datos', 'Pasos ejecutados', 'Resultado esperado', 'Resultado observado', 'Estado: aprobado/fallido/pendiente', 'Fecha y hora del computador', 'Nombre de captura', 'Incidencia y nueva ejecución']:
    p(label + ': [COMPLETAR]')
sub('Anexo E. Referencias')
p('ISO. ISO/IEC 25010:2023. Systems and software engineering — SQuaRE — Product quality model. https://www.iso.org/standard/78176.html. Consultado el 7 de octubre de 2026.')
p('OWASP Foundation. Application Security Verification Standard (ASVS). https://owasp.org/projects/asvs. Consultado el 7 de octubre de 2026.')
p('Repositorio COOPAECA. Código, configuración y esquema revisados: README.md; docs/PWA_CAPACITOR_ANDROID.md; public/index.php; app/Core/Database.php; app/Helpers/auth.php; app/Helpers/security.php; app/Helpers/crypto.php; app/Services; tests/Unit; capacitor.config.json; dashboard_base.sql. Evidencia primaria del proyecto local.')

for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.name = 'Arial'; run.font.size = Pt(12)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('COOPAECA | Producto académico 2 | ')
run.font.name = 'Arial'; run.font.size = Pt(12)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
base = OUT / 'PA2_Apellido1_Apellido2_Apellido3'
doc.save(base.with_suffix('.docx'))
# Word opens RTF documents; retain the explicit .rtf extension to avoid
# representing this interchange format as a legacy binary .doc file.
def esc(s):
    out=''
    for c in s:
        if c in '\\{}': out+='\\'+c
        elif ord(c)>127:
            n=ord(c); out+=f'\\u{n if n<32768 else n-65536}?'
        else: out+=c
    return out
rtf = ['{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}\\paperw11906\\paperh16838\\margl1417\\margr1417\\margt1417\\margb1417\\f0\\fs24 ']
for text,bold in records:
    if text=='PAGE': rtf.append('\\page ')
    else: rtf.append('\\pard\\sl360\\slmult1\\sa120 '+('\\b ' if bold else '')+esc(text)+('\\b0' if bold else '')+'\\par\n')
rtf.append('}')
base.with_suffix('.rtf').write_text(''.join(rtf),encoding='ascii')
(OUT/'LEEME.txt').write_text('Documento editable en Arial 12. Completar portada, marcadores y evidencias. Si la universidad exige .doc: abrir el .docx en Microsoft Word, Archivo > Guardar como > Documento de Word 97-2003 (*.doc). Cambiar Apellido1_Apellido2_Apellido3 por los apellidos reales. No se ha renombrado un formato distinto como .doc. Se incluye una alternativa RTF compatible con Word.\n',encoding='utf-8')
print(base.with_suffix('.docx'))
print('Párrafos:',len(doc.paragraphs),'Palabras:',sum(len(t.split()) for t,b in records))
